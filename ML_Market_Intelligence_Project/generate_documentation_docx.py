import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def build_docx(filename="Project_Documentation.docx"):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. Title
    p_title_num = doc.add_paragraph()
    r = p_title_num.add_run("1. Title")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(15, 23, 42)

    p_title = doc.add_paragraph()
    r_main = p_title.add_run("Failure Prediction AI – Startup & Project Risk Analyzer")
    r_main.bold = True
    r_main.font.size = Pt(14)
    r_main.font.color.rgb = RGBColor(30, 41, 59)

    # Banner Box Table
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_table.autofit = False
    banner_table.columns[0].width = Inches(6.8)
    cell = banner_table.cell(0, 0)
    set_cell_background(cell, "1E293B")
    
    bp = cell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br1 = bp.add_run("Failure Prediction AI\n")
    br1.bold = True
    br1.font.size = Pt(16)
    br1.font.color.rgb = RGBColor(255, 255, 255)
    
    br2 = bp.add_run("Startup & Project Risk Analyzer\n\n")
    br2.font.size = Pt(11)
    br2.font.color.rgb = RGBColor(203, 213, 225)

    br3 = bp.add_run("PREDICT RISKS • MITIGATE FAILURES • ENSURE SUCCESS")
    br3.bold = True
    br3.font.size = Pt(9)
    br3.font.color.rgb = RGBColor(255, 255, 255)

    doc.add_paragraph()

    # 2. Project Statement
    p2_num = doc.add_paragraph()
    r = p2_num.add_run("2. Project Statement")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)

    p2_txt = doc.add_paragraph(
        "Entrepreneurs, students, startups, and project teams often invest significant time, effort, and "
        "resources into ideas without fully understanding the risks that may impact success. Many startup "
        "ventures, business plans, hackathon solutions, and college projects fail due to poor market fit, "
        "intense competition, budget limitations, operational challenges, or inadequate planning. "
        "This project aims to develop Failure Prediction AI, an intelligent decision-support platform that "
        "analyses startup ideas, business proposals, college projects, and innovation concepts to predict "
        "potential failure risks and recommend mitigation strategies. By leveraging Large Language Models (LLMs), "
        "strategic reasoning, market intelligence, risk assessment frameworks, and agentic AI workflows, "
        "the system provides data-driven insights into business viability and project success probability. "
        "The solution enables users to make informed decisions, reduce uncertainty, and improve the likelihood "
        "of successful execution."
    )
    p2_txt.paragraph_format.line_spacing = 1.2
    p2_txt.paragraph_format.space_after = Pt(8)

    # 3. Outcomes
    p3_num = doc.add_paragraph()
    r = p3_num.add_run("3. Outcomes")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)

    outcomes = [
        "AI-powered evaluation of startup ideas, business plans, and project proposals.",
        "Automated risk scoring and success probability estimation.",
        "SWOT analysis generation for strategic assessment.",
        "Competitor and market opportunity analysis.",
        "Identification of critical business, technical, financial, and operational risks.",
        "Personalized mitigation strategies and improvement recommendations.",
        "Interactive dashboard for risk visualization and decision support.",
        "Comprehensive project viability and feasibility reports."
    ]
    for o in outcomes:
        bp = doc.add_paragraph(style='List Bullet')
        r = bp.add_run(o)
        r.font.size = Pt(10)
        bp.paragraph_format.space_after = Pt(3)

    # 4. Modules to be Implemented
    p4_num = doc.add_paragraph()
    r = p4_num.add_run("4. Modules to be Implemented")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)

    mods = [
        ("1. Idea & Project Information Collection", [
            "Collect startup, project, business plan, and hackathon idea details.",
            "Capture target market, budget, competition, resources, and objectives.",
            "Store proposal information for future analysis."
        ]),
        ("2. Market & Competitor Intelligence Engine", [
            "Analyse target market characteristics.",
            "Identify key competitors and industry challenges.",
            "Assess market opportunities and growth potential."
        ]),
        ("3. Risk Assessment & Scoring Engine", [
            "Evaluate business, financial, operational, and technical risks.",
            "Generate overall risk score and success probability.",
            "Prioritize high-impact risk factors."
        ]),
        ("4. SWOT Analysis & Strategic Evaluation", [
            "Generate Strengths, Weaknesses, Opportunities, and Threats analysis.",
            "Evaluate project feasibility and sustainability.",
            "Identify strategic advantages and gaps."
        ]),
        ("5. Recommendation & Mitigation Engine", [
            "Generate actionable recommendations.",
            "Suggest mitigation strategies for identified risks.",
            "Recommend execution roadmaps and improvement plans."
        ]),
        ("6. Dashboard & Reporting System", [
            "Interactive dashboard for risk visualization.",
            "Generate detailed feasibility and assessment reports.",
            "Provide business intelligence and strategic insights."
        ])
    ]
    for m_title, items in mods:
        mp = doc.add_paragraph()
        r = mp.add_run(m_title)
        r.bold = True
        r.font.size = Pt(11)
        for it in items:
            bp = doc.add_paragraph(style='List Bullet')
            bp.add_run(it)
            bp.paragraph_format.space_after = Pt(2)

    doc.add_page_break()

    # 5. Week-wise Implementation
    p5_num = doc.add_paragraph()
    r = p5_num.add_run("5. Week-wise Module Implementation & High-Level Requirements")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)

    milestones = [
        ("Milestone 1 (Weeks 1–2) – Data Collection & Market Intelligence", [
            "Develop startup and project submission forms.",
            "Implement market and competitor analysis modules.",
            "Store project information in PostgreSQL."
        ]),
        ("Milestone 2 (Weeks 3–4) – Risk Assessment & SWOT Analysis", [
            "Develop risk scoring engine.",
            "Generate SWOT analysis automatically.",
            "Implement project feasibility assessment."
        ]),
        ("Milestone 3 (Weeks 5–6) – Recommendations & Strategic Reasoning", [
            "Integrate Gemini/OpenAI for strategic recommendations.",
            "Develop mitigation and improvement suggestion engine.",
            "Implement LangGraph agent workflows."
        ]),
        ("Milestone 4 (Weeks 7–8) – Dashboard & Deployment", [
            "Build risk analytics dashboard.",
            "Generate comprehensive assessment reports.",
            "Perform testing, optimization, and deployment."
        ])
    ]
    for ms_title, items in milestones:
        msp = doc.add_paragraph()
        r = msp.add_run(ms_title)
        r.bold = True
        r.font.size = Pt(11)
        for it in items:
            bp = doc.add_paragraph(style='List Bullet')
            bp.add_run(it)
            bp.paragraph_format.space_after = Pt(2)

    # 6. Evaluation Criteria
    p6_num = doc.add_paragraph()
    r = p6_num.add_run("6. Evaluation Criteria")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)

    eval_crit = [
        ("Milestone 1 (Week 2)", [
            "Successfully collect and process startup/project information.",
            "Market and competitor analysis generated correctly."
        ]),
        ("Milestone 2 (Week 4)", [
            "Risk scoring model produces consistent evaluations.",
            "SWOT analysis generated successfully for submitted projects."
        ]),
        ("Milestone 3 (Week 6)", [
            "AI recommendations align with identified risks.",
            "Mitigation strategies generated successfully."
        ]),
        ("Milestone 4 (Week 8)", [
            "Dashboard fully operational.",
            "End-to-end project risk analysis workflow functional.",
            "Assessment reports generated within 5 seconds.",
            "User satisfaction score ≥85% during testing and evaluation."
        ])
    ]
    for ev_title, items in eval_crit:
        evp = doc.add_paragraph()
        r = evp.add_run(ev_title)
        r.bold = True
        r.font.size = Pt(11)
        for it in items:
            bp = doc.add_paragraph(style='List Bullet')
            bp.add_run(it)
            bp.paragraph_format.space_after = Pt(2)

    doc.add_page_break()

    # 7. Workflow Diagram Table
    p7_num = doc.add_paragraph()
    r = p7_num.add_run("7. Workflow Diagram")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)

    p7_sub = doc.add_paragraph()
    r = p7_sub.add_run("Failure Prediction AI – Startup & Project Risk Analyzer Workflow Flow")
    r.italic = True
    r.font.size = Pt(10)

    wf_steps = [
        ("1. Ingestion", "Startup & Project Submission", "Collects Title, Industry, Model, Budget, Description, Team, Resources."),
        ("2. Profiling", "Project Profiling Engine", "Structures inputs and validates parameters for downstream intelligence engines."),
        ("3. Market Intel", "Market & Competitor Analysis", "Computes TAM ($2.4B), SAM ($850M), SOM ($12M), and historical market trends."),
        ("4. Risk Scoring", "5-Vector Risk Engine", "Evaluates Market, Financial, Competition, Technical, and Operational Risks (1-5)."),
        ("5. Evaluation", "SWOT & Feasibility Engine", "Produces Strengths, Weaknesses, Opportunities, Threats, and Feasibility Score."),
        ("6. Reasoning", "LangGraph Agent Workflows", "Runs 6-node reasoning: Analyzes risks, generates action cards & mitigations."),
        ("7. Delivery", "Interactive Dashboard & Reports", "Visualizes KPI cards, area trend chart, and generates instant PDF/DOCX downloads.")
    ]
    t7 = doc.add_table(rows=1, cols=3)
    t7.alignment = WD_TABLE_ALIGNMENT.CENTER
    t7.columns[0].width = Inches(1.5)
    t7.columns[1].width = Inches(2.2)
    t7.columns[2].width = Inches(3.1)
    
    hdr = t7.rows[0].cells
    hdr[0].paragraphs[0].add_run("Phase").bold = True
    hdr[1].paragraphs[0].add_run("Module Component").bold = True
    hdr[2].paragraphs[0].add_run("Process & Output Details").bold = True
    for c in hdr:
        set_cell_background(c, "0F172A")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for p, c, d in wf_steps:
        row = t7.add_row().cells
        row[0].paragraphs[0].add_run(p)
        row[1].paragraphs[0].add_run(c)
        row[2].paragraphs[0].add_run(d)

    doc.add_paragraph()

    # 8. Architecture Diagram Table
    p8_num = doc.add_paragraph()
    r = p8_num.add_run("8. Architecture Diagram")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)

    arch_layers = [
        ("User Layer", "Entrepreneurs, Students, Founders, Analysts", "User submissions, queries, scenarios."),
        ("Presentation Layer", "Streamlit Cyberpunk UI, CSS3", "OLED dark theme, KPI cards, charts."),
        ("Application Layer", "Streamlit Controller (app.py)", "State management, validation, routing."),
        ("Market Intel Layer", "market_analysis.py, Groq API", "TAM/SAM/SOM sizing, competitor insights."),
        ("Risk Assessment Layer", "risk_engine.py, swot_analysis.py", "5-risk vector scoring, SWOT matrix derivation."),
        ("AI Orchestration Layer", "LangGraph, StateGraph, LangChain", "6-node multi-step sequential reasoning."),
        ("LLM Layer", "Groq Llama-3 70B / OpenAI / Gemini", "Zero-shot JSON reasoning, mitigations."),
        ("Data Layer", "Supabase (PostgreSQL)", "Projects, risks, SWOT, and logs persistence."),
        ("Reporting Layer", "ReportLab (PDF), python-docx (Word)", "On-demand report generation in < 5 seconds.")
    ]
    t8 = doc.add_table(rows=1, cols=3)
    t8.alignment = WD_TABLE_ALIGNMENT.CENTER
    t8.columns[0].width = Inches(1.8)
    t8.columns[1].width = Inches(2.2)
    t8.columns[2].width = Inches(2.8)
    
    hdr8 = t8.rows[0].cells
    hdr8[0].paragraphs[0].add_run("Architecture Layer").bold = True
    hdr8[1].paragraphs[0].add_run("Technologies / Modules").bold = True
    hdr8[2].paragraphs[0].add_run("Core Responsibilities").bold = True
    for c in hdr8:
        set_cell_background(c, "1E293B")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for l, m, resp in arch_layers:
        row = t8.add_row().cells
        row[0].paragraphs[0].add_run(l)
        row[1].paragraphs[0].add_run(m)
        row[2].paragraphs[0].add_run(resp)

    doc.add_page_break()

    # 9. Database Schema Table
    p9_num = doc.add_paragraph()
    r = p9_num.add_run("9. Database Schema")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)

    db_tables = [
        ("users", "user_id (UUID)", "-", "name, email, role, created_at"),
        ("projects", "project_id (BIGSERIAL)", "user_id", "project_name, industry, business_model, budget, target_market, description"),
        ("analytics_logs", "log_id (BIGSERIAL)", "user_id", "activity_type, timestamp"),
        ("competitor_analysis", "analysis_id (BIGSERIAL)", "project_id", "competitor_name, market_position, competitive_risk, analysis_date"),
        ("swot_analysis", "swot_id (BIGSERIAL)", "project_id", "strengths[], weaknesses[], opportunities[], threats[], generated_at"),
        ("risk_assessments", "risk_id (BIGSERIAL)", "project_id", "risk_category, risk_score (1-5), risk_description, priority_level"),
        ("success_predictions", "prediction_id (BIGSERIAL)", "project_id", "success_probability (%), overall_risk_score, prediction_date"),
        ("recommendations", "recommendation_id (BIGSERIAL)", "project_id", "recommendation_text, risk_mitigation, priority, generated_at"),
        ("assessment_reports", "report_id (BIGSERIAL)", "project_id", "report_path, generated_at"),
        ("analytics", "usg_id (BIGSERIAL)", "project_id", "project_type, budget, markets, created_at, generated_at")
    ]
    t9 = doc.add_table(rows=1, cols=4)
    t9.alignment = WD_TABLE_ALIGNMENT.CENTER
    t9.columns[0].width = Inches(1.5)
    t9.columns[1].width = Inches(1.4)
    t9.columns[2].width = Inches(1.0)
    t9.columns[3].width = Inches(2.9)
    
    hdr9 = t9.rows[0].cells
    hdr9[0].paragraphs[0].add_run("Table Name").bold = True
    hdr9[1].paragraphs[0].add_run("Primary Key").bold = True
    hdr9[2].paragraphs[0].add_run("Foreign Key").bold = True
    hdr9[3].paragraphs[0].add_run("Key Attributes").bold = True
    for c in hdr9:
        set_cell_background(c, "0F172A")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for t_name, pk, fk, attrs in db_tables:
        row = t9.add_row().cells
        row[0].paragraphs[0].add_run(t_name)
        row[1].paragraphs[0].add_run(pk)
        row[2].paragraphs[0].add_run(fk)
        row[3].paragraphs[0].add_run(attrs)

    doc.save(filename)
    print(f"Documentation DOCX successfully created: {filename}")

if __name__ == "__main__":
    build_docx("Project_Documentation.docx")
